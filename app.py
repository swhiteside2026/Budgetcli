import csv
import io
import json
import os
import re
from datetime import date, datetime
from functools import wraps

from flask import Flask, Response, flash, g, jsonify, redirect, render_template, request, session, url_for

from budgetcli import auth
from budgetcli.mailer import send_password_reset
from budgetcli.models import VALID_CATEGORIES, RecurringTransaction, Transaction
from budgetcli.reports import category_breakdown, monthly_summary, overall_balance
from budgetcli.storage import Storage

app = Flask(__name__)
app.secret_key = "budgetcli-dev-secret"

THEMES = [
    ("cupcake", "Pink",   "#ec4899"),
    ("dark",    "Dark",   "#1d232a"),
    ("nord",    "Blue",   "#5e81ac"),
    ("forest",  "Green",  "#1eb854"),
    ("dracula", "Purple", "#bd93f9"),
]
_VALID_THEMES = {t[0] for t in THEMES}

# Snapshot of the built-in categories so per-request custom category syncing
# can reset to the base list without cross-user contamination.
_BASE_CATEGORIES: list[str] = list(VALID_CATEGORIES)


# ── Auth helpers ──────────────────────────────────────────────────────────────

def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if "username" not in session:
            return redirect(url_for("login"))
        return f(*args, **kwargs)
    return decorated


@app.before_request
def _setup_request() -> None:
    if "username" in session:
        g.store = Storage(auth.user_data_path(session["username"]))
        # Reset to built-ins then add this user's custom categories so that
        # Transaction validation sees the right set for the current user.
        VALID_CATEGORIES[:] = _BASE_CATEGORIES
        for cat in g.store.load_custom_categories():
            if cat not in VALID_CATEGORIES:
                VALID_CATEGORIES.append(cat)
        # Apply stored theme preference on the first request of a new session.
        if "theme" not in session:
            stored = g.store.load_profile().get("default_theme", "")
            if stored in _VALID_THEMES:
                session["theme"] = stored


_EMAIL_RE = re.compile(r'^[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}$')


def _valid_email(email: str) -> bool:
    return bool(_EMAIL_RE.match(email))


def _valid_phone(phone: str) -> bool:
    digits = re.sub(r'\D', '', phone)
    return 7 <= len(digits) <= 15


def _check_and_notify_limits(
    store: Storage,
    transactions: list,
    limits: dict,
    year: int,
    month: int,
    near_threshold: float,
) -> list[dict]:
    """Create NEAR/OVER notifications for any budget category that has crossed a threshold.

    Returns the list of notification dicts that were actually inserted (empty when all
    were suppressed by the dedup logic or no threshold was crossed).
    """
    new_alerts: list[dict] = []
    spent = category_breakdown(transactions, year, month)
    for cat, limit in limits.items():
        if limit <= 0:
            continue
        cat_spent = spent.get(cat, 0.0)
        pct = cat_spent / limit * 100
        if cat_spent > limit:
            result = store.add_notification(
                f"{cat.capitalize()} is over budget — "
                f"${cat_spent:.2f} of ${limit:.2f} spent ({pct:.0f}%)",
                "over", cat,
            )
        elif cat_spent >= limit * near_threshold:
            result = store.add_notification(
                f"{cat.capitalize()} is approaching its limit — "
                f"${cat_spent:.2f} of ${limit:.2f} spent ({pct:.0f}%)",
                "near", cat,
            )
        else:
            result = None
        if result is not None:
            new_alerts.append({"id": result["id"], "message": result["message"], "type": result["type"]})
    return new_alerts


def _all_categories() -> list[str]:
    custom = g.store.load_custom_categories()
    return VALID_CATEGORIES + [c for c in custom if c not in VALID_CATEGORIES]


def _budget_data(
    transactions: list[Transaction],
    limits: dict[str, float],
    year: int,
    month: int,
    near_threshold: float = 0.8,
) -> list[dict]:
    spent = category_breakdown(transactions, year, month)
    result = []
    for cat, limit in limits.items():
        cat_spent = spent.get(cat, 0.0)
        pct = min(cat_spent / limit * 100, 100) if limit > 0 else 0
        over = cat_spent > limit
        near = not over and cat_spent >= limit * near_threshold
        result.append({
            "category": cat,
            "limit": limit,
            "spent": cat_spent,
            "pct": round(pct, 1),
            "bar_class": "progress-error" if over else ("progress-warning" if near else "progress-primary"),
            "label_class": "text-error font-bold" if over else "",
            "status": "OVER" if over else ("NEAR" if near else "OK"),
            "status_class": "badge-error" if over else ("badge-warning" if near else "badge-primary"),
        })
    return result


@app.context_processor
def inject_globals() -> dict:
    current_theme = session.get("theme", "cupcake")
    base: dict = {
        "current_theme": current_theme,
        "themes": THEMES,
        "current_user": session.get("username"),
    }
    if "username" not in session:
        return {**base, "global_balance": 0, "global_balance_abs": 0,
                "unread_notifications": 0, "budget_alerts": []}
    txns = g.store.load_transactions()
    bal = overall_balance(txns)
    unread = sum(1 for n in g.store.load_notifications() if not n.get("read"))
    # Pop any alerts queued by the previous POST so they show exactly once
    budget_alerts: list[dict] = []
    if "pending_budget_alerts" in session:
        budget_alerts = session.pop("pending_budget_alerts")
    return {**base, "global_balance": bal, "global_balance_abs": abs(bal),
            "unread_notifications": unread, "budget_alerts": budget_alerts}


# ── Auth routes ───────────────────────────────────────────────────────────────

@app.route("/login", methods=["GET", "POST"])
def login():
    if "username" in session:
        return redirect(url_for("dashboard"))
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        if auth.verify_user(username, password):
            session["username"] = username.lower()
            return redirect(url_for("dashboard"))
        flash("Invalid username or password.", "error")
    return render_template("login.html")


@app.route("/register", methods=["GET", "POST"])
def register():
    if "username" in session:
        return redirect(url_for("dashboard"))
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        confirm = request.form.get("confirm", "")
        email = request.form.get("email", "").strip()
        err = auth.validate_username(username)
        if err:
            flash(err, "error")
        elif len(password) < 8:
            flash("Password must be at least 8 characters.", "error")
        elif password != confirm:
            flash("Passwords do not match.", "error")
        elif email and not _valid_email(email):
            flash("Please enter a valid email address.", "error")
        elif auth.user_exists(username):
            flash("That username is already taken.", "error")
        else:
            auth.create_user(username, password)
            session["username"] = username.lower()
            if email:
                Storage(auth.user_data_path(username.lower())).save_profile({"email": email})
            flash("Account created! Welcome to BudgetBalancer.", "success")
            return redirect(url_for("dashboard"))
    return render_template("register.html")


@app.route("/logout", methods=["POST"])
def logout():
    session.pop("username", None)
    return redirect(url_for("login"))


@app.route("/forgot-password", methods=["GET", "POST"])
def forgot_password():
    if "username" in session:
        return redirect(url_for("dashboard"))
    if request.method == "POST":
        email = request.form.get("email", "").strip()
        username = auth.find_user_by_email(email)
        if username:
            token, expiry_iso = auth.generate_reset_token()
            Storage(auth.user_data_path(username)).save_reset_token(token, expiry_iso)
            app_url = os.environ.get("APP_URL", request.url_root.rstrip("/"))
            reset_url = f"{app_url}/reset-password?token={token}"
            try:
                send_password_reset(email, reset_url)
            except Exception:
                pass  # Never reveal send failure — message stays neutral
        # Always show the same message to avoid revealing registered emails
        flash(
            "If an account exists for that email, a reset link has been sent.",
            "info",
        )
        return redirect(url_for("forgot_password"))
    return render_template("forgot_password.html")


@app.route("/reset-password", methods=["GET", "POST"])
def reset_password():
    if "username" in session:
        return redirect(url_for("dashboard"))
    if request.method == "POST":
        token = request.form.get("token", "")
        username = auth.find_user_by_reset_token(token)
        if not username:
            flash("This reset link is invalid or has expired.", "error")
            return redirect(url_for("forgot_password"))
        new_pw = request.form.get("new_password", "")
        confirm_pw = request.form.get("confirm_password", "")
        if len(new_pw) < 8:
            return render_template(
                "reset_password.html", token=token,
                error="Password must be at least 8 characters.",
            )
        if new_pw != confirm_pw:
            return render_template(
                "reset_password.html", token=token,
                error="Passwords do not match.",
            )
        auth.change_password(username, new_pw)
        Storage(auth.user_data_path(username)).clear_reset_token()
        flash("Password reset successfully. You can now sign in.", "success")
        return redirect(url_for("login"))
    # GET: validate token before showing the form
    token = request.args.get("token", "")
    username = auth.find_user_by_reset_token(token)
    if not username:
        return render_template("reset_password.html", token=None, error="invalid")
    return render_template("reset_password.html", token=token, error=None)


# ── Theme route (no login required) ──────────────────────────────────────────

@app.route("/theme/set", methods=["POST"])
def set_theme():
    theme = request.form.get("theme", "cupcake")
    if theme in _VALID_THEMES:
        session["theme"] = theme
    return redirect(request.referrer or url_for("dashboard"))


# ── Dashboard ─────────────────────────────────────────────────────────────────

@app.route("/")
@login_required
def dashboard() -> str:
    today = date.today()
    transactions = g.store.load_transactions()
    limits = g.store.load_limits()
    income, expenses, savings = monthly_summary(transactions, today.year, today.month)
    balance = overall_balance(transactions)
    near_threshold = g.store.load_alert_threshold() / 100
    budget_data = _budget_data(transactions, limits, today.year, today.month, near_threshold)
    recent = sorted(transactions, key=lambda t: t.date, reverse=True)[:10]
    return render_template(
        "dashboard.html",
        transactions=recent,
        budget_data=budget_data,
        income=income,
        expenses=expenses,
        savings=savings,
        savings_abs=abs(savings),
        balance=balance,
        balance_abs=abs(balance),
        month_label=today.strftime("%B %Y"),
        total_count=len(transactions),
    )


# ── Transactions ──────────────────────────────────────────────────────────────

@app.route("/transactions")
@login_required
def transactions() -> str:
    all_txns = g.store.load_transactions()
    indexed = sorted(enumerate(all_txns), key=lambda x: x[1].date, reverse=True)
    return render_template(
        "transactions.html",
        indexed_transactions=indexed,
        categories=_all_categories(),
        today=date.today().isoformat(),
    )


@app.route("/transactions/add", methods=["POST"])
@login_required
def add_transaction_route():
    try:
        txn = Transaction(
            amount=float(request.form["amount"]),
            category=request.form["category"],
            date=request.form["date"],
            note=request.form.get("note", ""),
        )
        g.store.add_transaction(txn)
        today = date.today()
        new_alerts = _check_and_notify_limits(
            g.store, g.store.load_transactions(), g.store.load_limits(),
            today.year, today.month, g.store.load_alert_threshold() / 100,
        )
        if new_alerts:
            session["pending_budget_alerts"] = new_alerts
        flash("Transaction added.", "success")
    except (ValueError, KeyError) as e:
        flash(f"Error: {e}", "error")
    return redirect(url_for("transactions"))


@app.route("/transactions/<int:idx>/delete", methods=["POST"])
@login_required
def delete_transaction_route(idx: int):
    try:
        g.store.delete_transaction(idx)
        flash("Transaction deleted.", "success")
    except IndexError:
        flash("Transaction not found.", "error")
    return redirect(url_for("transactions"))


@app.route("/transactions/bulk-delete", methods=["POST"])
@login_required
def bulk_delete_transactions_route():
    raw_ids = request.form.getlist("ids")
    try:
        indices = sorted({int(i) for i in raw_ids if i.strip()}, reverse=True)
    except ValueError:
        flash("Invalid selection.", "error")
        return redirect(url_for("transactions"))
    if not indices:
        flash("No transactions selected.", "error")
        return redirect(url_for("transactions"))
    deleted = 0
    for idx in indices:
        try:
            g.store.delete_transaction(idx)
            deleted += 1
        except IndexError:
            pass
    flash(f"Deleted {deleted} transaction{'s' if deleted != 1 else ''}.", "success")
    return redirect(url_for("transactions"))


@app.route("/transactions/<int:idx>/edit", methods=["POST"])
@login_required
def edit_transaction_route(idx: int):
    try:
        txn = Transaction(
            amount=float(request.form["amount"]),
            category=request.form["category"],
            date=request.form["date"],
            note=request.form.get("note", ""),
        )
        g.store.update_transaction(idx, txn)
        flash("Transaction updated.", "success")
    except (ValueError, KeyError) as e:
        flash(f"Error: {e}", "error")
    return redirect(url_for("transactions"))


# ── Budget limits ─────────────────────────────────────────────────────────────

@app.route("/limits")
@login_required
def limits() -> str:
    today = date.today()
    current_limits = g.store.load_limits()
    transactions = g.store.load_transactions()
    near_threshold = g.store.load_alert_threshold() / 100
    budget_data = _budget_data(transactions, current_limits, today.year, today.month, near_threshold)
    available = [c for c in _all_categories() if c != "income"]
    return render_template(
        "limits.html",
        budget_data=budget_data,
        current_limits=current_limits,
        available_categories=available,
        month_label=today.strftime("%B %Y"),
    )


@app.route("/limits/set", methods=["POST"])
@login_required
def set_limit_route():
    try:
        category = request.form["category"]
        amount = float(request.form["amount"])
        if amount <= 0:
            raise ValueError("Amount must be greater than zero")
        g.store.set_limit(category, amount)
        flash(f"Limit set: {category} → ${amount:.2f}/month.", "success")
    except (ValueError, KeyError) as e:
        flash(f"Error: {e}", "error")
    return redirect(url_for("limits"))


@app.route("/limits/<category>/delete", methods=["POST"])
@login_required
def delete_limit_route(category: str):
    g.store.remove_limit(category)
    flash(f"Limit removed for {category}.", "success")
    return redirect(url_for("limits"))


# ── Settings ──────────────────────────────────────────────────────────────────

@app.route("/settings")
@login_required
def settings() -> str:
    return render_template(
        "settings.html",
        alert_threshold=g.store.load_alert_threshold(),
        profile=g.store.load_profile(),
    )


@app.route("/settings/profile", methods=["POST"])
@login_required
def save_profile_route():
    display_name = request.form.get("display_name", "").strip()
    email = request.form.get("email", "").strip()
    phone = request.form.get("phone", "").strip()
    if email and not _valid_email(email):
        flash("Please enter a valid email address.", "error")
        return redirect(url_for("settings"))
    if phone and not _valid_phone(phone):
        flash("Please enter a valid phone number (7–15 digits).", "error")
        return redirect(url_for("settings"))
    g.store.save_profile({"display_name": display_name, "email": email, "phone": phone})
    flash("Profile updated.", "success")
    return redirect(url_for("settings"))


@app.route("/settings/theme-preference", methods=["POST"])
@login_required
def save_theme_preference_route():
    theme = request.form.get("theme", "cupcake")
    if theme not in _VALID_THEMES:
        flash("Invalid theme selection.", "error")
        return redirect(url_for("settings"))
    session["theme"] = theme
    g.store.save_profile({"default_theme": theme})
    flash("Default theme updated.", "success")
    return redirect(url_for("settings"))


@app.route("/settings/password", methods=["POST"])
@login_required
def change_password_route():
    username = session["username"]
    current_pw = request.form.get("current_password", "")
    new_pw = request.form.get("new_password", "")
    confirm_pw = request.form.get("confirm_password", "")
    if not auth.verify_user(username, current_pw):
        flash("Current password is incorrect.", "error")
        return redirect(url_for("settings"))
    if len(new_pw) < 8:
        flash("New password must be at least 8 characters.", "error")
        return redirect(url_for("settings"))
    if new_pw != confirm_pw:
        flash("New passwords do not match.", "error")
        return redirect(url_for("settings"))
    auth.change_password(username, new_pw)
    flash("Password changed successfully.", "success")
    return redirect(url_for("settings"))


@app.route("/settings/alert-threshold", methods=["POST"])
@login_required
def save_alert_threshold_route():
    try:
        value = int(request.form["alert_threshold"])
        if not 1 <= value <= 99:
            raise ValueError("Threshold must be between 1 and 99.")
        g.store.save_alert_threshold(value)
        flash(f"Alert threshold updated to {value}%.", "success")
    except (ValueError, KeyError) as e:
        flash(f"Error: {e}", "error")
    return redirect(url_for("settings"))


# ── Notifications ────────────────────────────────────────────────────────────

@app.route("/notifications")
@login_required
def notifications() -> str:
    notifs = g.store.load_notifications()
    g.store.mark_all_read()
    return render_template("notifications.html", notifications=notifs)


@app.route("/notifications/<notif_id>/delete", methods=["POST"])
@login_required
def delete_notification_route(notif_id: str):
    g.store.delete_notification(notif_id)
    return redirect(url_for("notifications"))


@app.route("/notifications/clear", methods=["POST"])
@login_required
def clear_notifications_route():
    g.store.save_notifications([])
    return redirect(url_for("notifications"))


# ── Recurring ─────────────────────────────────────────────────────────────────

@app.route("/recurring")
@login_required
def recurring() -> str:
    today = date.today()
    rec_list = g.store.load_recurring()
    indexed = [(i, r, r.is_due(today)) for i, r in enumerate(rec_list)]
    due_count = sum(1 for _, _, due in indexed if due)
    return render_template(
        "recurring.html",
        indexed_recurring=indexed,
        categories=_all_categories(),
        today=today,
        due_count=due_count,
    )


@app.route("/recurring/add", methods=["POST"])
@login_required
def add_recurring_route():
    try:
        rt = RecurringTransaction(
            amount=float(request.form["amount"]),
            category=request.form["category"],
            frequency="monthly",
            note=request.form.get("note", ""),
        )
        g.store.add_recurring(rt)
        flash("Recurring transaction added.", "success")
    except (ValueError, KeyError) as e:
        flash(f"Error: {e}", "error")
    return redirect(url_for("recurring"))


@app.route("/recurring/apply", methods=["POST"])
@login_required
def apply_recurring_route():
    today = date.today()
    rec_list = g.store.load_recurring()
    applied = 0
    for rt in rec_list:
        if rt.is_due(today):
            g.store.add_transaction(
                Transaction(amount=rt.amount, category=rt.category, date=today, note=rt.note)
            )
            rt.last_applied = today
            applied += 1
    g.store.save_recurring(rec_list)
    if applied:
        new_alerts = _check_and_notify_limits(
            g.store, g.store.load_transactions(), g.store.load_limits(),
            today.year, today.month, g.store.load_alert_threshold() / 100,
        )
        if new_alerts:
            session["pending_budget_alerts"] = new_alerts
    label = "success" if applied else "info"
    plural = "s" if applied != 1 else ""
    flash(f"Applied {applied} recurring transaction{plural}.", label)
    return redirect(url_for("recurring"))


@app.route("/recurring/<int:idx>/delete", methods=["POST"])
@login_required
def delete_recurring_route(idx: int):
    rec_list = g.store.load_recurring()
    if 0 <= idx < len(rec_list):
        rec_list.pop(idx)
        g.store.save_recurring(rec_list)
        flash("Recurring transaction removed.", "success")
    else:
        flash("Not found.", "error")
    return redirect(url_for("recurring"))


# ── Export ────────────────────────────────────────────────────────────────────

@app.route("/export")
@login_required
def export() -> str:
    today = date.today()
    return render_template(
        "export.html",
        default_from=date(today.year, today.month, 1).isoformat(),
        default_to=today.isoformat(),
    )


@app.route("/export/download")
@login_required
def export_download():
    from_str = request.args.get("from_date", "")
    to_str = request.args.get("to_date", "")
    try:
        from_date = date.fromisoformat(from_str) if from_str else None
        to_date = date.fromisoformat(to_str) if to_str else None
    except ValueError:
        flash("Invalid date format.", "error")
        return redirect(url_for("export"))

    txns = g.store.load_transactions()
    if from_date:
        txns = [t for t in txns if t.date >= from_date]
    if to_date:
        txns = [t for t in txns if t.date <= to_date]
    txns = sorted(txns, key=lambda t: t.date)

    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["date", "category", "amount", "note"])
    for t in txns:
        writer.writerow([t.date.isoformat(), t.category, round(t.amount, 2), t.note])

    return Response(
        buf.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": "attachment; filename=transactions.csv"},
    )


# ── Categories ────────────────────────────────────────────────────────────────

@app.route("/categories")
@login_required
def categories() -> str:
    custom = g.store.load_custom_categories()
    txns = g.store.load_transactions()
    rec = g.store.load_recurring()
    used: set[str] = {t.category for t in txns} | {r.category for r in rec}
    return render_template(
        "categories.html",
        default_categories=VALID_CATEGORIES,
        custom_categories=custom,
        used_categories=used,
    )


@app.route("/categories/add", methods=["POST"])
@login_required
def add_category_route():
    raw = request.form.get("name", "").strip().lower().replace(" ", "-")
    name = re.sub(r"[^a-z0-9-]", "", raw)
    if len(name) < 2 or len(name) > 30 or name.startswith("-") or name.endswith("-"):
        flash("Category name must be 2–30 characters, letters/numbers/hyphens, no leading or trailing hyphens.", "error")
        return redirect(url_for("categories"))
    if name in _all_categories():
        flash(f'Category "{name}" already exists.', "error")
        return redirect(url_for("categories"))
    g.store.add_custom_category(name)
    VALID_CATEGORIES.append(name)
    flash(f'Category "{name}" added.', "success")
    return redirect(url_for("categories"))


@app.route("/categories/<name>/delete", methods=["POST"])
@login_required
def delete_category_route(name: str):
    custom = g.store.load_custom_categories()
    if name not in custom:
        flash(f'"{name}" is a built-in category and cannot be deleted.', "error")
        return redirect(url_for("categories"))
    g.store.remove_custom_category(name)
    if name in VALID_CATEGORIES:
        VALID_CATEGORIES.remove(name)
    flash(f'Category "{name}" removed.', "success")
    return redirect(url_for("categories"))


# ── Excel Import ──────────────────────────────────────────────────────────────

def _parse_import_date(s: str, prefer_dmy: bool = False) -> date:
    # ISO / year-first formats are always unambiguous
    for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    # Slash/dash-separated: order determined by caller's format preference
    ordered = (
        ["%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y", "%m-%d-%Y"]
        if prefer_dmy
        else ["%m/%d/%Y", "%m-%d-%Y", "%d/%m/%Y", "%d-%m-%Y"]
    )
    for fmt in ordered:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Cannot parse date: {s!r}")


def _analyse_date_column(values: list[str]) -> tuple[str, bool]:
    """Scan raw date strings and infer whether the file uses MM/DD or DD/MM.

    Returns (fmt, needs_user_choice).
    fmt is 'mdy' (MM/DD/YYYY) or 'dmy' (DD/MM/YYYY).
    needs_user_choice is True when auto-detection cannot resolve the format.
    """
    us_only = eu_only = uncertain = 0
    for raw in values:
        s = raw.strip()
        if not s:
            continue
        # Year-first formats are unambiguous — don't affect the count
        iso = False
        for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
            try:
                datetime.strptime(s, fmt)
                iso = True
                break
            except ValueError:
                pass
        if iso:
            continue
        ok_us = ok_eu = False
        for fmt in ("%m/%d/%Y", "%m-%d-%Y"):
            try:
                datetime.strptime(s, fmt)
                ok_us = True
                break
            except ValueError:
                pass
        for fmt in ("%d/%m/%Y", "%d-%m-%Y"):
            try:
                datetime.strptime(s, fmt)
                ok_eu = True
                break
            except ValueError:
                pass
        if ok_us and ok_eu:
            uncertain += 1
        elif ok_us:
            us_only += 1
        elif ok_eu:
            eu_only += 1

    # A single unambiguous value settles the whole column
    if eu_only > 0 and us_only == 0:
        return "dmy", False
    if us_only > 0 and eu_only == 0:
        return "mdy", False
    # Conflict or no diagnostic values — need user input
    if uncertain > 0 or (us_only > 0 and eu_only > 0):
        return "mdy", True   # default to mdy; user must confirm
    return "mdy", False


@app.route("/import")
@login_required
def import_excel():
    return render_template("import.html", step="upload", categories=_all_categories())


@app.route("/import/upload", methods=["POST"])
@login_required
def import_upload():
    file = request.files.get("file")
    if not file or not file.filename:
        flash("No file selected.", "error")
        return redirect(url_for("import_excel"))

    fname = file.filename.lower()
    if not (fname.endswith(".xlsx") or fname.endswith(".xls")):
        flash("Only .xlsx and .xls files are accepted.", "error")
        return redirect(url_for("import_excel"))

    try:
        if fname.endswith(".xlsx"):
            import openpyxl

            wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
            ws = wb.active

            def _to_str(v) -> str:
                if v is None:
                    return ""
                if isinstance(v, datetime):
                    return v.date().isoformat()
                if isinstance(v, date):
                    return v.isoformat()
                return str(v).strip()

            string_rows = [[_to_str(c) for c in row] for row in ws.iter_rows(values_only=True)]
        else:
            import xlrd

            wb = xlrd.open_workbook(file_contents=file.read())
            ws = wb.sheet_by_index(0)

            def _xlrd_str(cell) -> str:
                if cell.ctype == xlrd.XL_CELL_DATE:
                    return xlrd.xldate_as_datetime(cell.value, wb.datemode).date().isoformat()
                if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    return ""
                if cell.ctype == xlrd.XL_CELL_NUMBER:
                    v = cell.value
                    return str(int(v)) if v == int(v) else str(v)
                return str(cell.value).strip()

            string_rows = [
                [_xlrd_str(ws.cell(i, j)) for j in range(ws.ncols)]
                for i in range(ws.nrows)
            ]
    except Exception as e:
        flash(f"Could not read Excel file: {e}", "error")
        return redirect(url_for("import_excel"))

    if not string_rows:
        flash("The file is empty.", "error")
        return redirect(url_for("import_excel"))

    headers = string_rows[0]
    data_rows = [r for r in string_rows[1:] if any(c for c in r)]

    if not data_rows:
        flash("No data rows found (the file only has a header row).", "error")
        return redirect(url_for("import_excel"))

    _ALIASES: dict[str, list[str]] = {
        "date":     ["date", "transaction date", "trans date", "txn date", "trans. date"],
        "amount":   ["amount", "value", "sum", "total", "price", "debit", "credit"],
        "category": ["category", "cat", "type", "transaction type", "txn type"],
        "note":     ["note", "notes", "description", "desc", "memo", "comment", "details", "narration"],
    }
    detected: dict[str, int] = {}
    for field, aliases in _ALIASES.items():
        for i, h in enumerate(headers):
            if h.lower().strip() in aliases and field not in detected:
                detected[field] = i

    # Analyse the date column so the preview step can flag ambiguous formats
    date_col_idx = detected.get("date")
    if date_col_idx is not None:
        date_values = [
            row[date_col_idx]
            for row in data_rows
            if date_col_idx < len(row)
        ]
        date_fmt, date_ambiguous = _analyse_date_column(date_values)
    else:
        date_fmt, date_ambiguous = "mdy", False

    temp_path = auth.user_data_path(session["username"]).parent / "import_temp.json"
    temp_path.write_text(
        json.dumps({
            "headers": headers,
            "rows": data_rows,
            "detected": detected,
            "date_fmt": date_fmt,
            "date_ambiguous": date_ambiguous,
        }),
        encoding="utf-8",
    )

    return redirect(url_for("import_preview"))


@app.route("/import/preview")
@login_required
def import_preview():
    temp_path = auth.user_data_path(session["username"]).parent / "import_temp.json"
    if not temp_path.exists():
        flash("No import in progress. Please upload a file first.", "error")
        return redirect(url_for("import_excel"))

    data = json.loads(temp_path.read_text(encoding="utf-8"))
    return render_template(
        "import.html",
        step="preview",
        headers=data["headers"],
        preview_rows=data["rows"][:50],
        total_rows=len(data["rows"]),
        detected=data["detected"],
        categories=_all_categories(),
        date_fmt=data.get("date_fmt", "mdy"),
        date_ambiguous=data.get("date_ambiguous", False),
    )


@app.route("/import/confirm", methods=["POST"])
@login_required
def import_confirm():
    temp_path = auth.user_data_path(session["username"]).parent / "import_temp.json"
    if not temp_path.exists():
        flash("No import in progress.", "error")
        return redirect(url_for("import_excel"))

    data = json.loads(temp_path.read_text(encoding="utf-8"))
    rows = data["rows"]

    mapping: dict[str, int] = {}
    for field in ("date", "amount", "category", "note"):
        val = request.form.get(f"col_{field}", "-1")
        try:
            col_idx = int(val)
            if col_idx >= 0:
                mapping[field] = col_idx
        except ValueError:
            pass

    if "date" not in mapping or "amount" not in mapping:
        flash("You must map both the Date and Amount columns.", "error")
        return redirect(url_for("import_preview"))

    prefer_dmy = request.form.get("date_fmt", "mdy") == "dmy"
    default_category = request.form.get("default_category", "other")
    all_cats = _all_categories()
    if default_category not in all_cats:
        default_category = "other"

    saved = 0
    skipped = 0
    for row in rows:
        try:
            raw_date = row[mapping["date"]].strip() if mapping["date"] < len(row) else ""
            raw_amount = (
                row[mapping["amount"]].strip().lstrip("$").replace(",", "")
                if mapping["amount"] < len(row) else ""
            )
            raw_cat = (
                row[mapping["category"]].strip().lower()
                if "category" in mapping and mapping["category"] < len(row)
                else ""
            )
            raw_note = (
                row[mapping["note"]].strip()
                if "note" in mapping and mapping["note"] < len(row)
                else ""
            )

            if not raw_date or not raw_amount:
                skipped += 1
                continue

            txn_date = _parse_import_date(raw_date, prefer_dmy=prefer_dmy)
            amount = abs(float(raw_amount))
            if amount == 0.0:
                skipped += 1
                continue

            category = raw_cat if raw_cat in all_cats else default_category
            g.store.add_transaction(
                Transaction(amount=amount, category=category, date=txn_date, note=raw_note[:200])
            )
            saved += 1
        except (ValueError, IndexError):
            skipped += 1

    temp_path.unlink(missing_ok=True)

    if saved:
        _today = date.today()
        new_alerts = _check_and_notify_limits(
            g.store, g.store.load_transactions(), g.store.load_limits(),
            _today.year, _today.month, g.store.load_alert_threshold() / 100,
        )
        if new_alerts:
            session["pending_budget_alerts"] = new_alerts
        msg = f"Imported {saved} transaction{'s' if saved != 1 else ''}."
        if skipped:
            msg += f" {skipped} row{'s' if skipped != 1 else ''} skipped (invalid or empty)."
        flash(msg, "success")
    else:
        flash(
            "No transactions could be imported. Check that your file has valid date and amount columns.",
            "error",
        )

    return redirect(url_for("transactions"))


# ── Voice API ─────────────────────────────────────────────────────────────────

@app.route("/api/voice-context")
@login_required
def voice_context_api():
    today = date.today()
    transactions = g.store.load_transactions()
    limits = g.store.load_limits()
    income, expenses, savings = monthly_summary(transactions, today.year, today.month)
    near_threshold = g.store.load_alert_threshold() / 100
    budget_data = _budget_data(transactions, limits, today.year, today.month, near_threshold)
    return jsonify({
        "categories": _all_categories(),
        "month_label": today.strftime("%B %Y"),
        "income": round(income, 2),
        "expenses": round(expenses, 2),
        "savings": round(savings, 2),
        "budget_data": [
            {
                "category": item["category"],
                "spent": round(item["spent"], 2),
                "limit": round(item["limit"], 2),
                "status": item["status"],
                "pct": item["pct"],
            }
            for item in budget_data
        ],
    })


@app.route("/api/voice-add", methods=["POST"])
@login_required
def voice_add_api():
    try:
        data = request.get_json(force=True)
        txn = Transaction(
            amount=float(data["amount"]),
            category=str(data["category"]),
            date=date.today().isoformat(),
            note="added by voice",
        )
        g.store.add_transaction(txn)
        today = date.today()
        new_alerts = _check_and_notify_limits(
            g.store, g.store.load_transactions(), g.store.load_limits(),
            today.year, today.month, g.store.load_alert_threshold() / 100,
        )
        unread_count = sum(1 for n in g.store.load_notifications() if not n.get("read"))
        return jsonify({"ok": True, "new_alerts": new_alerts, "unread_count": unread_count})
    except (ValueError, KeyError, TypeError) as e:
        return jsonify({"ok": False, "error": str(e)}), 400


if __name__ == "__main__":
    import os
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=False)
